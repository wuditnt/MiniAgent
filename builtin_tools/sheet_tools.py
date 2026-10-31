# pip install openpyxl python-calamine
"""
表格工具组（.xlsx / .xlsm / .xls / .xlsb / .ods）：读取与写入。

独立边界：
  · 只处理电子表格；纯文本 CSV / JSON 归 data_tools，分析聚合归 dataframe_tools。
  · 不与任何兄弟模块互相 import；公共能力只依赖 core。
  · 第三方库懒加载：缺 calamine 自动回退 openpyxl；缺库时只让当前工具返回提示。
"""

from __future__ import annotations

import csv
import io
import os
import re
import unicodedata

from core import read_path, safe_path, tool, workspace

TOOL_INFO = {
    "name": "表格",
    "version": "1.2",
    "author": "mini-agent",
    "desc": "读写 Excel（.xlsx/.xlsm/.xls/.xlsb/.ods），合并、拆分、CSV 互转、格式化、批量公式",
}


def _need(lib: str, pip_name: str = "") -> str:
    """统一的「缺第三方库」提示。"""
    return (
        f"[error] 需要第三方库 {lib}，当前环境未安装。\n"
        f"  安装：python -m pip install {pip_name or lib}\n"
        f"  （装到本项目虚拟环境，勿用系统 Python）"
    )


def _md_table(rows: list[list[str]]) -> str:
    """把二维数据渲染成 Markdown 表格（行宽按最长行对齐）。"""
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    head = rows[0] + [""] * (ncols - len(rows[0]))
    out = ["| " + " | ".join(head) + " |", "|" + "|".join(["---"] * ncols) + "|"]
    for r in rows[1:]:
        cells = [c for c in r] + [""] * (ncols - len(r))
        out.append("| " + " | ".join(c[:80].replace("\n", " ") for c in cells) + " |")
    return "\n".join(out)


# 只有 calamine 能读的旧格式 / 非 OOXML 格式，openpyxl 读不了
_CALAMINE_ONLY_EXTS = (".xls", ".xlsb", ".ods")


def _xlsx_cell(v) -> str:
    """单元格值转字符串。calamine 把数字统一给成 float，整数值去掉多余的 .0。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ---------------------------------------------------------------- 通用读写

#: 工作表名不允许的字符（Excel 硬限制）
_BAD_SHEET_CHARS = re.compile(r"[\[\]\*\?/\\:]")
#: 文件名里需要替换的字符
_BAD_FILE_CHARS = re.compile(r"[^\w.-]+")


def _sheet_name(raw: str, fallback: str = "Sheet") -> str:
    """把任意字符串变成合法工作表名（去非法字符、截 31 字符）。"""
    name = _BAD_SHEET_CHARS.sub("_", str(raw or "")).strip().strip("'") or fallback
    return name[:31]


def _unique_name(taken: set[str], raw: str, fallback: str = "Sheet") -> str:
    """在已有表名集合里生成不重名的表名（Excel 表名大小写不敏感）。"""
    base = _sheet_name(raw, fallback)
    name = base
    i = 2
    while name.lower() in taken:
        suffix = f"_{i}"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    taken.add(name.lower())
    return name


def _safe_stem(path: str) -> str:
    """文件主干安全化，避免输出名带路径分隔符。"""
    stem = os.path.splitext(os.path.basename(path))[0] or "sheet"
    return _BAD_FILE_CHARS.sub("_", stem).strip("._") or "sheet"


def _sheet_names(ap: str) -> list[str]:
    """列出全部工作表名。calamine 优先，openpyxl 兜底。"""
    try:
        from python_calamine import CalamineWorkbook

        return list(CalamineWorkbook.from_path(ap).sheet_names)
    except ImportError:
        pass
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"解析 {os.path.splitext(ap)[1]} 失败（{type(e).__name__}: {e}）") from e

    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("需要第三方库 openpyxl：python -m pip install openpyxl") from e
    try:
        wb = openpyxl.load_workbook(ap, read_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"解析 {os.path.splitext(ap)[1]} 失败（{type(e).__name__}: {e}）") from e


def _sheet_rows(ap: str, sheet: str = "", max_rows: int = 0, max_cols: int = 0) -> list[list]:
    """读取指定工作表（缺省第一张）的原始值。calamine 优先，openpyxl 兜底。"""
    try:
        from python_calamine import CalamineWorkbook

        wb = CalamineWorkbook.from_path(ap)
        names = list(wb.sheet_names)
        if not names:
            return []
        target = sheet if sheet in names else names[0]
        ws = wb.get_sheet_by_name(target)
        nrows = int(ws.height or 0)
        ncols = int(ws.width or 0)
        if max_rows > 0:
            nrows = min(nrows, max_rows)
        if max_cols > 0:
            ncols = min(ncols, max_cols)
        if nrows <= 0:
            return []
        rows = ws.to_python(nrows=nrows)
        return [[("" if c is None else c) for c in r[:ncols]] for r in rows]
    except ImportError:
        pass
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"解析 {os.path.splitext(ap)[1]} 失败（{type(e).__name__}: {e}）") from e

    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("需要第三方库 openpyxl：python -m pip install openpyxl") from e
    try:
        wb = openpyxl.load_workbook(ap, read_only=True, data_only=True)
        names = list(wb.sheetnames)
        if not names:
            wb.close()
            return []
        target = sheet if sheet in names else names[0]
        ws = wb[target]
        nrows = ws.max_row or 0
        ncols = ws.max_column or 0
        if max_rows > 0:
            nrows = min(nrows, max_rows)
        if max_cols > 0:
            ncols = min(ncols, max_cols)
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=nrows, max_col=ncols, values_only=True)] if nrows else []
        wb.close()
        return rows
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"解析 {os.path.splitext(ap)[1]} 失败（{type(e).__name__}: {e}）") from e


def _write_book(ap: str, sheets: list[tuple[str, list[list]]]) -> int:
    """把若干 (表名, 行) 写成一个 .xlsx，返回写入的行数。"""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    total = 0
    for title, rows in sheets:
        ws = wb.create_sheet(title=_sheet_name(title))
        for r in rows or []:
            ws.append(list(r) if isinstance(r, (list, tuple)) else [r])
            total += 1
    os.makedirs(os.path.dirname(ap) or workspace(), exist_ok=True)
    wb.save(ap)
    return total


def _csv_rows(ap: str, delimiter: str = "", encoding: str = "") -> list[list[str]]:
    """读 CSV：自动试 UTF-8-BOM / GBK，分隔符嗅探失败退回逗号。"""
    encodings = [encoding] if encoding else ["utf-8-sig", "gbk", "utf-8"]
    text = ""
    last: Exception | None = None
    for enc in encodings:
        try:
            with open(ap, "r", encoding=enc, newline="") as f:
                text = f.read()
            break
        except (UnicodeDecodeError, LookupError) as e:
            last = e
    else:
        raise ValueError(f"无法解码 CSV（尝试 {encodings}）：{last}")

    sep = delimiter
    if not sep:
        try:
            sep = csv.Sniffer().sniff(text[:2000], delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=sep)]


def _coerce(value: str):
    """把纯数字字符串转成数值；带前导零的编号（如「001」）保持文本。"""
    if not isinstance(value, str):
        return value
    s = value.strip()
    if not s:
        return ""
    digits = s.lstrip("-")
    if re.fullmatch(r"-?\d+", s) and not (len(digits) > 1 and digits.startswith("0")):
        try:
            return int(s)
        except ValueError:
            return value
    if re.fullmatch(r"-?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            return value
    return value


def _disp_width(value) -> int:
    """估算显示宽度：东亚宽字符按 2 计，避免中文列过窄。"""
    text = _xlsx_cell(value)
    width = 0
    for ch in text[:60]:
        width += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return width


def _xlsx_head(ap: str, names: list[str], target: str, nrows: int, ncols: int,
               shown_rows: int, shown_cols: int, asked: str) -> str:
    head = f"[{ap}] 工作表 {names} / 当前: {target} / {nrows} 行 × {ncols} 列"
    if asked and asked != target:
        head += f"（未找到表 {asked!r}，已回退到 {target}）"
    if nrows > shown_rows or ncols > shown_cols:
        head += f"\n…（仅显示前 {shown_rows} 行 × {shown_cols} 列）"
    return head


def _read_xlsx_calamine(ap: str, sheet: str, max_rows: int, max_cols: int) -> str:
    """python-calamine（Rust）：读取快，支持 xlsx / xlsm / xls / xlsb / ods。"""
    from python_calamine import CalamineWorkbook

    try:
        wb = CalamineWorkbook.from_path(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 解析 {os.path.splitext(ap)[1]} 失败（{type(e).__name__}: {e}）"

    names = list(wb.sheet_names)
    if not names:
        return f"[error] {ap} 没有任何工作表"
    target = sheet if sheet in names else names[0]

    ws = wb.get_sheet_by_name(target)
    nrows, ncols = int(ws.height or 0), int(ws.width or 0)
    shown_rows = min(nrows, max(1, max_rows))
    shown_cols = min(ncols, max(1, max_cols))

    data: list[list[str]] = []
    if shown_rows > 0:
        rows = ws.to_python(nrows=shown_rows)
        data = [[_xlsx_cell(c) for c in r[:shown_cols]] for r in rows]

    head = _xlsx_head(ap, names, target, nrows, ncols, shown_rows, shown_cols, sheet)
    return f"{head}\n{_md_table(data) if data else '（空表）'}"


def _read_xlsx_openpyxl(ap: str, sheet: str, max_rows: int, max_cols: int) -> str:
    """openpyxl 兜底（纯 Python，不支持 .xls/.xlsb/.ods）。"""
    try:
        import openpyxl
    except ImportError:
        return _need("openpyxl")

    try:
        wb = openpyxl.load_workbook(ap, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return f"[error] 解析 .xlsx 失败（{type(e).__name__}: {e}）"

    names = list(wb.sheetnames)
    if not names:
        return f"[error] {ap} 没有任何工作表"
    target = sheet if sheet in names else names[0]

    ws = wb[target]
    nrows = ws.max_row or 0
    ncols = ws.max_column or 0
    shown_rows = min(nrows, max(1, max_rows))
    shown_cols = min(ncols, max(1, max_cols))

    data = [
        ["" if v is None else str(v) for v in r]
        for r in ws.iter_rows(min_row=1, max_row=shown_rows, max_col=shown_cols, values_only=True)
    ]
    wb.close()

    head = _xlsx_head(ap, names, target, nrows, ncols, shown_rows, shown_cols, sheet)
    return f"{head}\n{_md_table(data) if data else '（空表）'}"


@tool(
    "read_xlsx",
    "读取 Excel 工作表并以 Markdown 表格返回，支持 .xlsx/.xlsm/.xls/.xlsb/.ods。"
    "engine=auto 时优先用 python-calamine（更快，且能读旧版 .xls），缺库自动回退 openpyxl。"
    "不给 sheet 时列出全部表名并展示第一张。",
    group="表格",
    params={
        "path": {"type": "string", "description": "Excel 文件路径"},
        "sheet": {"type": "string", "description": "工作表名，留空取第一张", "default": ""},
        "max_rows": {"type": "integer", "description": "最多显示行数", "default": 50},
        "max_cols": {"type": "integer", "description": "最多显示列数", "default": 30},
        "engine": {"type": "string", "enum": ["auto", "calamine", "openpyxl"],
                   "description": "读取引擎", "default": "auto"},
    },
    required=["path"],
)
def read_xlsx(path: str, sheet: str = "", max_rows: int = 50, max_cols: int = 30,
              engine: str = "auto") -> str:
    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"

    engine = (engine or "auto").strip().lower()
    ext = os.path.splitext(ap)[1].lower()

    if engine == "openpyxl":
        if ext in _CALAMINE_ONLY_EXTS:
            return f"[error] openpyxl 不支持 {ext}，请改用 engine=calamine"
        return _read_xlsx_openpyxl(ap, sheet, max_rows, max_cols)

    try:
        import python_calamine  # noqa: F401
    except ImportError:
        if engine == "calamine" or ext in _CALAMINE_ONLY_EXTS:
            return _need("python-calamine", "python-calamine")
        return _read_xlsx_openpyxl(ap, sheet, max_rows, max_cols)

    return _read_xlsx_calamine(ap, sheet, max_rows, max_cols)


@tool(
    "write_xlsx",
    "把二维数组写成 Excel(.xlsx)。rows 是行数组，每行是单元格数组。"
    "依赖第三方库 openpyxl。",
    group="表格",
    params={
        "path": {"type": "string", "description": "要写入的 xlsx 路径"},
        "rows": {"type": "array", "items": {"type": "array"},
                 "description": '二维数组，如 [["姓名","年龄"],["张三",18]]'},
        "sheet": {"type": "string", "description": "工作表名", "default": "Sheet1"},
    },
    required=["path", "rows"],
)
def write_xlsx(path: str, rows: list, sheet: str = "Sheet1") -> str:
    try:
        import openpyxl
    except ImportError:
        return _need("openpyxl")

    ap = safe_path(path)
    os.makedirs(os.path.dirname(ap) or workspace(), exist_ok=True)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet or "Sheet1"
        for r in rows or []:
            ws.append(list(r) if isinstance(r, (list, tuple)) else [r])
        wb.save(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 生成 .xlsx 失败（{type(e).__name__}: {e}）"

    return f"[ok] 已写入 {ap}（表 {ws.title}，{ws.max_row} 行 × {ws.max_column} 列，{os.path.getsize(ap)} 字节）"


# ---------------------------------------------------------------- 合并 / 拆分


@tool(
    "xlsx_merge",
    "把多个 Excel 文件合并成一个 .xlsx。mode=sheets 时逐个文件、逐张工作表追加为不同工作表"
    "（文件名会加进表名避免重名）；mode=rows 时把各文件指定工作表的数据行纵向拼接，"
    "has_header=True 只保留首个文件的表头。读取支持 .xlsx/.xlsm/.xls/.xlsb/.ods。",
    group="表格",
    params={
        "paths": {"type": "array", "items": {"type": "string"},
                  "description": "要合并的 Excel 文件列表"},
        "out": {"type": "string", "description": "输出 .xlsx 路径"},
        "mode": {"type": "string", "enum": ["sheets", "rows"],
                 "description": "sheets=按工作表合并 / rows=按数据行拼接", "default": "sheets"},
        "sheet": {"type": "string", "description": "mode=rows 时读取的工作表名，留空取第一张", "default": ""},
        "has_header": {"type": "boolean", "description": "mode=rows 时除首个文件外跳过表头", "default": True},
        "source_column": {"type": "string",
                          "description": "mode=rows 时追加一列来源文件名，值作为列名；留空不加", "default": ""},
        "out_sheet": {"type": "string", "description": "mode=rows 时输出工作表名", "default": "Sheet1"},
    },
    required=["paths", "out"],
)
def xlsx_merge(paths: list, out: str, mode: str = "sheets", sheet: str = "",
               has_header: bool = True, source_column: str = "", out_sheet: str = "Sheet1") -> str:
    try:
        import openpyxl
    except ImportError:
        return _need("openpyxl")

    files = [str(p).strip() for p in (paths or []) if str(p).strip()]
    if not files:
        return "[error] paths 不能为空"

    mode = (mode or "sheets").strip().lower()
    if mode not in ("sheets", "rows"):
        return f"[error] mode 只能是 sheets 或 rows，收到 {mode!r}"

    out_ap = safe_path(out)
    if not out_ap.lower().endswith(".xlsx"):
        return "[error] 输出必须是 .xlsx 路径"

    for p in files:
        if not os.path.isfile(read_path(p)):
            return f"[error] 文件不存在：{p}"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    taken: set[str] = set()

    if mode == "sheets":
        for p in files:
            ap = read_path(p)
            try:
                names = _sheet_names(ap)
            except ValueError as e:
                return f"[error] {e}"
            if not names:
                continue
            stem = _safe_stem(p)
            for n in names:
                raw = n if len(files) == 1 else f"{stem}_{n}"
                title = _unique_name(taken, raw, "Sheet")
                try:
                    rows = _sheet_rows(ap, n)
                except ValueError as e:
                    return f"[error] {e}"
                ws = wb.create_sheet(title=title)
                for r in rows:
                    ws.append(list(r))
        if not wb.sheetnames:
            return "[error] 源文件没有任何工作表"
        os.makedirs(os.path.dirname(out_ap) or workspace(), exist_ok=True)
        wb.save(out_ap)
        return (f"[ok] 已合并 {len(files)} 个文件 → {out_ap}"
                f"（{len(wb.sheetnames)} 张表：{wb.sheetnames}）")

    merged: list[list] = []
    for i, p in enumerate(files):
        ap = read_path(p)
        try:
            rows = _sheet_rows(ap, sheet)
        except ValueError as e:
            return f"[error] {e}"
        if not rows:
            continue
        if i == 0:
            if source_column:
                block = [[source_column] + list(rows[0])]
                block += [[os.path.basename(p)] + list(r) for r in rows[1:]]
            else:
                block = rows
        else:
            body = rows[1:] if has_header else rows
            block = [[os.path.basename(p)] + list(r) for r in body] if source_column else body
        merged.extend(block)

    ws = wb.create_sheet(title=_sheet_name(out_sheet, "Sheet1"))
    for r in merged:
        ws.append(list(r))
    os.makedirs(os.path.dirname(out_ap) or workspace(), exist_ok=True)
    wb.save(out_ap)
    return f"[ok] 已拼接 {len(files)} 个文件 → {out_ap}（{ws.title}，{len(merged)} 行 × {ws.max_column} 列）"


@tool(
    "xlsx_split",
    "拆分 Excel。mode=sheets 时每张工作表存成独立 .xlsx；mode=rows 时按 rows_per_file 行切成多个文件"
    "（has_header=True 时每个分片都保留表头）。输出到 out_dir，文件名自动带源名与表名/分片号。",
    group="表格",
    params={
        "path": {"type": "string", "description": "源 Excel 文件路径"},
        "out_dir": {"type": "string", "description": "输出目录"},
        "mode": {"type": "string", "enum": ["sheets", "rows"],
                 "description": "sheets=按工作表拆分 / rows=按行数拆分", "default": "sheets"},
        "rows_per_file": {"type": "integer", "description": "mode=rows 时每个文件的数据行数", "default": 1000},
        "has_header": {"type": "boolean", "description": "mode=rows 时每个分片保留表头", "default": True},
        "sheet": {"type": "string", "description": "mode=rows 时读取的工作表名，留空取第一张", "default": ""},
    },
    required=["path", "out_dir"],
)
def xlsx_split(path: str, out_dir: str, mode: str = "sheets", rows_per_file: int = 1000,
               has_header: bool = True, sheet: str = "") -> str:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return _need("openpyxl")

    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"

    mode = (mode or "sheets").strip().lower()
    if mode not in ("sheets", "rows"):
        return f"[error] mode 只能是 sheets 或 rows，收到 {mode!r}"

    out_ap = safe_path(out_dir)
    os.makedirs(out_ap, exist_ok=True)
    stem = _safe_stem(path)
    written: list[str] = []

    if mode == "sheets":
        try:
            names = _sheet_names(ap)
        except ValueError as e:
            return f"[error] {e}"
        if not names:
            return f"[error] {ap} 没有任何工作表"

        taken: set[str] = set()
        for n in names:
            try:
                rows = _sheet_rows(ap, n)
            except ValueError as e:
                return f"[error] {e}"
            title = _unique_name(taken, n, "Sheet")
            target = os.path.join(out_ap, f"{stem}_{title}.xlsx")
            _write_book(target, [(title, rows)])
            written.append(f"{title} → {os.path.basename(target)}（{len(rows)} 行）")
        return f"[ok] 已按工作表拆成 {len(written)} 个文件 → {out_ap}\n  " + "\n  ".join(written)

    try:
        rows = _sheet_rows(ap, sheet)
    except ValueError as e:
        return f"[error] {e}"
    if not rows:
        return f"[error] {ap} 工作表为空"
    header = [list(rows[0])] if has_header else []
    body = rows[1:] if has_header else rows
    if not body:
        return f"[error] {ap} 只有表头，没有数据行可拆分"

    size = max(1, int(rows_per_file or 1))
    title = _sheet_name(sheet, "Sheet1")
    for i in range(0, len(body), size):
        chunk = header + [list(r) for r in body[i:i + size]]
        idx = i // size + 1
        target = os.path.join(out_ap, f"{stem}_part{idx}.xlsx")
        _write_book(target, [(title, chunk)])
        written.append(f"part{idx} → {os.path.basename(target)}（{len(chunk)} 行）")
    return (f"[ok] 已按 {size} 行/份拆成 {len(written)} 个文件 → {out_ap}"
            f"（共 {len(body)} 条数据行）\n  " + "\n  ".join(written))


# ---------------------------------------------------------------- CSV 互转


@tool(
    "xlsx_to_csv",
    "把 Excel 工作表导出为 CSV（默认 UTF-8-BOM，双击即可在 Excel 里正确显示中文）。"
    "sheet 留空导出全部工作表，每张表一个 csv。读取支持 .xlsx/.xlsm/.xls/.xlsb/.ods。",
    group="表格",
    params={
        "path": {"type": "string", "description": "源 Excel 文件路径"},
        "out_dir": {"type": "string", "description": "输出目录"},
        "sheet": {"type": "string", "description": "只导出这张表；留空导出全部", "default": ""},
        "delimiter": {"type": "string", "description": "分隔符，默认英文逗号", "default": ","},
        "max_rows": {"type": "integer", "description": "每张表最多导出多少行，0 表示全部", "default": 0},
    },
    required=["path", "out_dir"],
)
def xlsx_to_csv(path: str, out_dir: str, sheet: str = "", delimiter: str = ",",
                max_rows: int = 0) -> str:
    sep = delimiter or ","
    if len(sep) != 1:
        return f"[error] delimiter 必须是单个字符，收到 {delimiter!r}"

    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"
    try:
        names = _sheet_names(ap)
    except ValueError as e:
        return f"[error] {e}"
    if not names:
        return f"[error] {ap} 没有任何工作表"
    if sheet and sheet not in names:
        return f"[error] 未找到工作表 {sheet!r}，可用：{names}"

    out_ap = safe_path(out_dir)
    os.makedirs(out_ap, exist_ok=True)
    stem = _safe_stem(path)
    taken: set[str] = set()
    written: list[str] = []

    for n in ([sheet] if sheet else names):
        try:
            rows = _sheet_rows(ap, n, max_rows=max_rows)
        except ValueError as e:
            return f"[error] {e}"
        title = _unique_name(taken, n, "Sheet")
        target = os.path.join(out_ap, f"{stem}_{title}.csv")
        with open(target, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=sep)
            for r in rows:
                writer.writerow([_xlsx_cell(c) for c in r])
        written.append(f"表 {title} → {os.path.basename(target)}（{len(rows)} 行）")

    return f"[ok] 已导出 {len(written)} 个 CSV → {out_ap}\n  " + "\n  ".join(written)


@tool(
    "csv_to_xlsx",
    "把 CSV 汇总成一个 Excel：每个 CSV 一张工作表（表名取文件名）。"
    "自动识别 UTF-8 / GBK 与逗号/分号/Tab 分隔符；coerce_numbers=True 时把纯数字转成数值"
    "（带前导零的编号保持文本）。",
    group="表格",
    params={
        "paths": {"type": "array", "items": {"type": "string"}, "description": "CSV 文件列表"},
        "out": {"type": "string", "description": "输出 .xlsx 路径"},
        "delimiter": {"type": "string", "description": "分隔符；留空自动识别", "default": ""},
        "encoding": {"type": "string", "description": "编码；留空自动试 UTF-8/GBK", "default": ""},
        "coerce_numbers": {"type": "boolean", "description": "把纯数字文本转成数值", "default": True},
        "sheet": {"type": "string", "description": "只有一个 CSV 时指定工作表名", "default": ""},
    },
    required=["paths", "out"],
)
def csv_to_xlsx(paths: list, out: str, delimiter: str = "", encoding: str = "",
                coerce_numbers: bool = True, sheet: str = "") -> str:
    try:
        import openpyxl
    except ImportError:
        return _need("openpyxl")

    files = [str(p).strip() for p in (paths or []) if str(p).strip()]
    if not files:
        return "[error] paths 不能为空"

    out_ap = safe_path(out)
    if not out_ap.lower().endswith(".xlsx"):
        return "[error] 输出必须是 .xlsx 路径"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    taken: set[str] = set()
    total = 0

    for i, p in enumerate(files):
        ap = read_path(p)
        if not os.path.isfile(ap):
            return f"[error] 文件不存在：{p}"
        try:
            rows = _csv_rows(ap, delimiter, encoding)
        except ValueError as e:
            return f"[error] {e}"

        raw_title = sheet if (len(files) == 1 and sheet) else _safe_stem(p)
        title = _unique_name(taken, raw_title, f"CSV{i + 1}")
        ws = wb.create_sheet(title=title)
        for r in rows:
            ws.append([_coerce(c) if coerce_numbers else c for c in r])
            total += 1

    if not wb.sheetnames:
        return "[error] 没有可写入的 CSV 内容"
    os.makedirs(os.path.dirname(out_ap) or workspace(), exist_ok=True)
    wb.save(out_ap)
    return (f"[ok] 已把 {len(files)} 个 CSV 写成 {out_ap}"
            f"（{len(wb.sheetnames)} 张表，{total} 行：{wb.sheetnames}）")


# ---------------------------------------------------------------- 格式 / 公式


@tool(
    "xlsx_format",
    "格式化 Excel 工作表：表头加粗 + 底色 + 居中、冻结表头、自动筛选、按内容自适应列宽。"
    "只支持写回 .xlsx/.xlsm（.xls/.xlsb/.ods 请先另存为 xlsx）。依赖第三方库 openpyxl。",
    group="表格",
    params={
        "path": {"type": "string", "description": "要格式化的 Excel 路径（原地覆盖）"},
        "sheet": {"type": "string", "description": "工作表名，留空取当前活动表", "default": ""},
        "header_row": {"type": "integer", "description": "表头所在行号；0 表示不设置表头样式", "default": 1},
        "header_bold": {"type": "boolean", "description": "表头加粗", "default": True},
        "header_fill": {"type": "string", "description": "表头底色（十六进制，如 DDEBF7）；留空不填色", "default": "DDEBF7"},
        "freeze_header": {"type": "boolean", "description": "冻结表头（向下滚动时固定）", "default": True},
        "auto_filter": {"type": "boolean", "description": "给表头加自动筛选", "default": True},
        "auto_width": {"type": "boolean", "description": "按内容自适应列宽（中文按双宽计算）", "default": True},
        "max_col_width": {"type": "integer", "description": "列宽上限（字符数）", "default": 60},
    },
    required=["path"],
)
def xlsx_format(path: str, sheet: str = "", header_row: int = 1, header_bold: bool = True,
                header_fill: str = "DDEBF7", freeze_header: bool = True, auto_filter: bool = True,
                auto_width: bool = True, max_col_width: int = 60) -> str:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return _need("openpyxl")

    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"
    ext = os.path.splitext(ap)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        return f"[error] openpyxl 不能写回 {ext}，请先另存为 .xlsx"

    try:
        wb = openpyxl.load_workbook(ap, keep_vba=(ext == ".xlsm"))
    except Exception as e:  # noqa: BLE001
        return f"[error] 打开 {ext} 失败（{type(e).__name__}: {e}）"
    if sheet and sheet not in wb.sheetnames:
        return f"[error] 未找到工作表 {sheet!r}，可用：{wb.sheetnames}"
    ws = wb[sheet] if sheet else wb.active

    done: list[str] = []
    hrow = max(0, int(header_row or 0))
    if hrow and ws.max_row >= hrow:
        fill = PatternFill("solid", fgColor=header_fill.lstrip("#")) if header_fill else None
        for cell in ws[hrow]:
            if header_bold:
                font = Font(bold=True, name=cell.font.name, size=cell.font.size, color=cell.font.color)
                cell.font = font
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        done.append("表头样式")

        if freeze_header:
            ws.freeze_panes = f"A{hrow + 1}"
            done.append("冻结表头")
        if auto_filter and ws.max_row > hrow:
            ws.auto_filter.ref = ws.dimensions
            done.append("自动筛选")

    if auto_width and ws.max_column:
        cap = max(8, int(max_col_width or 60))
        for col in range(1, ws.max_column + 1):
            widest = max((_disp_width(ws.cell(row=r, column=col).value)
                          for r in range(1, ws.max_row + 1)), default=0)
            ws.column_dimensions[get_column_letter(col)].width = min(cap, max(8, widest + 2))
        done.append("自适应列宽")

    try:
        wb.save(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 保存失败（{type(e).__name__}: {e}）"

    ops = " / ".join(done) if done else "无（所有样式已关闭）"
    return f"[ok] 已格式化 {ws.title}（{ops}）→ {ap}"


@tool(
    "xlsx_formula",
    "向单元格区域批量写入 Excel 公式：formula 写在区域左上角，其余单元格按相对引用自动平移"
    '（如 cell_range="C2:C20"、formula="=A2*B2" 会得到 C3=A3*B3 …）。'
    "openpyxl 只写公式不做计算，Excel / WPS / LibreOffice 打开时自动求值。"
    "只支持 .xlsx/.xlsm。依赖第三方库 openpyxl。",
    group="表格",
    params={
        "path": {"type": "string", "description": "要写入公式的 Excel 路径（原地覆盖）"},
        "cell_range": {"type": "string", "description": '目标区域，如 "C2:C20" 或 "C2:F10"'},
        "formula": {"type": "string", "description": '公式，必须以 = 开头，如 "=A2*B2" 或 "=SUM(A2:A10)"'},
        "sheet": {"type": "string", "description": "工作表名，留空取当前活动表", "default": ""},
    },
    required=["path", "cell_range", "formula"],
)
def xlsx_formula(path: str, cell_range: str, formula: str, sheet: str = "") -> str:
    try:
        import openpyxl
        from openpyxl.formula.translate import Translator
        from openpyxl.utils import get_column_letter, range_boundaries
    except ImportError:
        return _need("openpyxl")

    text = (formula or "").strip()
    if not text.startswith("="):
        return f"[error] formula 必须以 = 开头，收到 {formula!r}"

    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"
    ext = os.path.splitext(ap)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        return f"[error] openpyxl 不能写回 {ext}，请先另存为 .xlsx"

    try:
        bounds = range_boundaries((cell_range or "").strip().upper())
    except ValueError:
        return f"[error] cell_range 不合法：{cell_range!r}"
    min_col, min_row, max_col, max_row = bounds

    try:
        wb = openpyxl.load_workbook(ap, keep_vba=(ext == ".xlsm"))
    except Exception as e:  # noqa: BLE001
        return f"[error] 打开 {ext} 失败（{type(e).__name__}: {e}）"
    if sheet and sheet not in wb.sheetnames:
        return f"[error] 未找到工作表 {sheet!r}，可用：{wb.sheetnames}"
    ws = wb[sheet] if sheet else wb.active

    origin = f"{get_column_letter(min_col)}{min_row}"
    translator = Translator(text, origin=origin)
    count = 0
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            coord = f"{get_column_letter(col)}{row}"
            ws[coord] = translator.translate_formula(coord)
            count += 1

    try:
        wb.save(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 保存失败（{type(e).__name__}: {e}）"

    return (f"[ok] 已写入 {count} 个公式到 {ws.title}!{cell_range}"
            f"（起点 {origin}，打开 Excel 时自动计算）")


