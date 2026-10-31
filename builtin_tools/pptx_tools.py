# pip install python-pptx
"""
PPT 演示工具组（.pptx）：读取每页文字/表格/备注、按结构化提纲生成演示文稿。

独立边界：
  · 只处理 .pptx；不解析 .docx / .xlsx / .pdf。
  · 不与任何兄弟模块互相 import；公共能力只依赖 core。
  · 第三方库懒加载：缺 python-pptx 时只让当前工具返回可读提示。
"""

from __future__ import annotations

import os

from core import clip, read_path, safe_path, tool, workspace

TOOL_INFO = {
    "name": "PPT 演示",
    "version": "1.0",
    "author": "mini-agent",
    "desc": "读取 PPTX 每页文字/表格/备注，按结构化提纲生成演示文稿（依赖 python-pptx）",
}

DEFAULT_MAX_CHARS = 8000


def _need(lib: str, pip_name: str = "") -> str:
    """统一的「缺第三方库」提示。"""
    return (
        f"[error] 需要第三方库 {lib}，当前环境未安装。\n"
        f"  安装：python -m pip install {pip_name or lib}\n"
        f"  （装到本项目虚拟环境，勿用系统 Python）"
    )


def _md_table(rows: list[list[str]]) -> str:
    """把二维数据渲染成 Markdown 表格（行宽按最长行对齐）。"""
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    head = rows[0] + [""] * (ncols - len(rows[0]))
    out = ["| " + " | ".join(head) + " |", "|" + "|".join(["---"] * ncols) + "|"]
    for r in rows[1:]:
        cells = [c for c in r] + [""] * (ncols - len(r))
        out.append("| " + " | ".join(c[:80].replace("\n", " ") for c in cells) + " |")
    return "\n".join(out)


def _slide_title(slide) -> str:
    try:
        title = slide.shapes.title
        return (title.text or "").strip() if title is not None else ""
    except Exception:  # noqa: BLE001
        return ""


def _slide_body(slide) -> list[str]:
    """收集正文段落与表格；标题形状单独处理，避免重复。"""
    out: list[str] = []
    title_shape = None
    try:
        title_shape = slide.shapes.title
    except Exception:  # noqa: BLE001
        pass

    for shape in slide.shapes:
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue
        if getattr(shape, "has_table", False):
            rows = [[("" if c.text is None else c.text).strip() for c in r.cells]
                    for r in shape.table.rows]
            if rows:
                out.append(_md_table(rows))
            continue
        if getattr(shape, "has_text_frame", False):
            for p in shape.text_frame.paragraphs:
                text = (p.text or "").strip()
                if text:
                    out.append("  " * int(getattr(p, "level", 0) or 0) + "- " + text)
    return out


@tool(
    "read_pptx",
    "读取 PowerPoint(.pptx)：按页输出标题、正文要点、表格与备注，转成 Markdown。"
    "只读取文字与表格结构，不导出图片。依赖第三方库 python-pptx。",
    group="文档·PPT",
    params={
        "path": {"type": "string", "description": "pptx 文件路径"},
        "max_slides": {"type": "integer", "description": "最多读取页数", "default": 50},
        "max_chars": {"type": "integer", "description": "返回最大字符数", "default": 8000},
        "include_notes": {"type": "boolean", "description": "是否包含演讲者备注", "default": True},
    },
    required=["path"],
)
def read_pptx(path: str, max_slides: int = 50, max_chars: int = DEFAULT_MAX_CHARS,
              include_notes: bool = True) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return _need("python-pptx")

    ap = read_path(path)
    if not os.path.isfile(ap):
        return f"[error] 文件不存在：{ap}"

    try:
        prs = Presentation(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 解析 .pptx 失败（{type(e).__name__}: {e}）"

    total = len(prs.slides)
    limit = max(1, max_slides)
    chunks: list[str] = []
    for i, slide in enumerate(prs.slides):
        if i >= limit:
            break
        title = _slide_title(slide)
        body = _slide_body(slide)
        chunk = [f"## 第 {i + 1} 页" + (f" · {title}" if title else "")]
        if body:
            chunk += body
        elif not title:
            chunk += ["- （本页无文字）"]
        if include_notes:
            try:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            except Exception:  # noqa: BLE001 —— 没有备注页时忽略
                notes = ""
            if notes:
                chunk.append(f"> 备注：{notes.replace(chr(10), chr(10) + '> ')}")
        chunks.append("\n".join(chunk))

    body = "\n\n".join(chunks)
    head = f"[{ap}] 共 {total} 页，读取前 {min(total, limit)} 页"
    if total > limit:
        head += f"（可用 max_slides 调整，当前上限 {limit}）"
    return clip(f"{head}\n{body}", max_chars)


def _set_bullets(placeholder, bullets: list) -> None:
    """把字符串列表写进正文占位符；支持 [{"text": ..., "level": n}] 形式。"""
    tf = placeholder.text_frame
    first = True
    for item in bullets:
        if isinstance(item, dict):
            text = str(item.get("text", "") or "")
            level = int(item.get("level", 0) or 0)
        else:
            text = str(item)
            level = 0
        if not text:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = text
        p.level = max(0, min(level, 8))
        first = False


@tool(
    "write_pptx",
    "按结构化提纲生成 PowerPoint(.pptx)。slides 是页数组，每页含 title、bullets、notes；"
    "可选 title 生成封面页。只生成标题+要点+备注，不做图表与复杂排版。"
    "输出文件已存在时直接覆盖。依赖第三方库 python-pptx。",
    group="文档·PPT",
    params={
        "path": {"type": "string", "description": "要写入的 pptx 路径"},
        "title": {"type": "string", "description": "封面标题，留空则不生成封面", "default": ""},
        "slides": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "本页标题"},
                    "bullets": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "本页要点，一行一条",
                    },
                    "notes": {"type": "string", "description": "演讲者备注"},
                },
                "required": ["title"],
            },
            "description": '页数组，如 [{"title":"背景","bullets":["现状","目标"]}]',
        },
    },
    required=["path", "slides"],
)
def write_pptx(path: str, slides: list, title: str = "") -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return _need("python-pptx")

    ap = safe_path(path)
    os.makedirs(os.path.dirname(ap) or workspace(), exist_ok=True)

    try:
        prs = Presentation()
        count = 0

        if title:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
            count += 1

        for item in slides or []:
            if not isinstance(item, dict):
                continue
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide_title = str(item.get("title", "") or "").strip()
            if slide.shapes.title is not None:
                slide.shapes.title.text = slide_title
            placeholder = None
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 1:
                    placeholder = ph
                    break
            if placeholder is None:
                placeholder = slide.placeholders[1]
            _set_bullets(placeholder, item.get("bullets") or [])
            notes = str(item.get("notes", "") or "")
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            count += 1

        prs.save(ap)
    except Exception as e:  # noqa: BLE001
        return f"[error] 生成 .pptx 失败（{type(e).__name__}: {e}）"

    return f"[ok] 已写入 {ap}（{count} 页，{os.path.getsize(ap)} 字节）"
